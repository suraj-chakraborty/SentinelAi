"""
sentinel/automation/deep_rpa.py
───────────────────────────────
Deep Autonomous Robotic Process Automation (RPA).
Provides window tiling/snapping, automated spreadsheet editing, and file organisation.
"""

import csv
import logging
import os
import shutil
import sys
from typing import Any

logger = logging.getLogger("DeepRPA")

try:
    import openpyxl

    _OPENPYXL_AVAILABLE = True
except ImportError:
    _OPENPYXL_AVAILABLE = False


class DeepRPA:
    """
    Autonomous multi-application RPA coordinator.
    """

    def snap_active_window(self, position: str = "left") -> bool:
        """
        Snaps the active foreground window to 'left', 'right', 'center', or 'maximize'.
        """
        if sys.platform != "win32":
            return False

        try:
            import ctypes

            user32 = ctypes.windll.user32

            hwnd = user32.GetForegroundWindow()
            if not hwnd:
                return False

            sw = user32.GetSystemMetrics(0)  # SM_CXSCREEN
            sh = user32.GetSystemMetrics(1)  # SM_CYSCREEN
            pos = position.lower().strip()

            if pos == "left":
                user32.ShowWindow(hwnd, 9)  # SW_RESTORE
                user32.MoveWindow(hwnd, 0, 0, sw // 2, sh, True)
            elif pos == "right":
                user32.ShowWindow(hwnd, 9)
                user32.MoveWindow(hwnd, sw // 2, 0, sw // 2, sh, True)
            elif pos == "center":
                user32.ShowWindow(hwnd, 9)
                w = int(sw * 0.7)
                h = int(sh * 0.8)
                x = (sw - w) // 2
                y = (sh - h) // 2
                user32.MoveWindow(hwnd, x, y, w, h, True)
            elif pos == "maximize":
                user32.ShowWindow(hwnd, 3)  # SW_MAXIMIZE
            else:
                return False

            logger.info(f"Snapped window {hwnd} to {pos}.")
            return True
        except Exception as e:
            logger.debug(f"Failed to snap window: {e}")
            return False

    def append_spreadsheet_row(self, filepath: str, row_data: list[Any], sheet_name: str | None = None) -> bool:
        """
        Appends a row to an Excel .xlsx or .csv spreadsheet.
        """
        try:
            ext = os.path.splitext(filepath)[1].lower()

            if ext == ".xlsx" and _OPENPYXL_AVAILABLE:
                if os.path.exists(filepath):
                    wb = openpyxl.load_workbook(filepath)
                else:
                    wb = openpyxl.Workbook()
                ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
                ws.append(row_data)
                wb.save(filepath)
                return True
            else:
                # Default CSV append
                csv_path = filepath if ext == ".csv" else f"{os.path.splitext(filepath)[0]}.csv"
                with open(csv_path, mode="a", newline="", encoding="utf-8") as f:
                    writer = csv.writer(f)
                    writer.writerow(row_data)
                return True
        except Exception as e:
            logger.error(f"Spreadsheet row append error: {e}")
            return False

    def organize_directory(self, target_dir: str) -> dict[str, Any]:
        """
        Autonomously sorts files into categorized subfolders by extension.
        """
        if not os.path.exists(target_dir):
            return {"success": False, "message": f"Directory not found: {target_dir}"}

        categories = {
            "Documents": [".pdf", ".docx", ".doc", ".txt", ".xlsx", ".pptx", ".csv"],
            "Images": [".jpg", ".jpeg", ".png", ".gif", ".webp", ".svg"],
            "Installers": [".exe", ".msi", ".zip", ".tar", ".gz", ".7z"],
            "Media": [".mp3", ".wav", ".mp4", ".mkv", ".mov"],
        }

        moved_count = 0
        try:
            for item in os.listdir(target_dir):
                item_path = os.path.join(target_dir, item)
                if os.path.isdir(item_path):
                    continue

                _, ext = os.path.splitext(item)
                ext = ext.lower()

                target_folder = None
                for cat, extensions in categories.items():
                    if ext in extensions:
                        target_folder = cat
                        break

                if target_folder:
                    dest_dir = os.path.join(target_dir, target_folder)
                    os.makedirs(dest_dir, exist_ok=True)
                    shutil.move(item_path, os.path.join(dest_dir, item))
                    moved_count += 1

            return {"success": True, "files_organized": moved_count}
        except Exception as e:
            logger.error(f"Directory organization error: {e}")
            return {"success": False, "error": str(e)}
